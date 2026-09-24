import csv
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from io import BytesIO, StringIO

from django.db.models import Prefetch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.clientes.models import Cliente
from apps.mantenimientos.models import MantenimientoRealizado, TipoMantenimiento
from apps.motos.models import Moto
from apps.notificaciones.models import (
    EventoSeguimientoMantenimiento,
    SeguimientoMantenimiento,
)
from apps.servicios.models import Servicio

from .utils import momento_local, valor_csv, valor_excel


FORMATO_FECHA = "DD/MM/YYYY"
FORMATO_FECHA_HORA = "DD/MM/YYYY HH:MM"
FORMATO_DINERO = "$ #,##0.00"
FORMATO_ENTERO = "#,##0"


@dataclass(frozen=True, slots=True)
class ColumnaExportacion:
    titulo: str
    ancho: int = 18
    formato: str | None = None
    envolver: bool = False


@dataclass(frozen=True, slots=True)
class TablaExportacion:
    slug: str
    hoja: str
    columnas: tuple[ColumnaExportacion, ...]
    obtener_filas: Callable[[], Iterable[tuple]]

    @property
    def encabezados(self):
        return tuple(columna.titulo for columna in self.columnas)


def _nombre_usuario(usuario):
    return usuario.get_username() if usuario is not None else ""


def filas_clientes():
    clientes = Cliente.objects.order_by("apellido", "nombre", "pk")
    for cliente in clientes:
        yield (
            cliente.pk,
            cliente.nombre,
            cliente.apellido,
            cliente.telefono,
            cliente.email,
            cliente.direccion,
            cliente.observaciones,
            cliente.activo,
            cliente.creado_en,
            cliente.actualizado_en,
        )


def filas_motos():
    motos = (
        Moto.objects.select_related("cliente")
        .only(
            "pk",
            "patente",
            "marca",
            "modelo",
            "anio",
            "cilindrada_cc",
            "color",
            "kilometraje_actual",
            "cliente_id",
            "cliente__nombre",
            "cliente__apellido",
            "cliente__telefono",
            "numero_chasis",
            "numero_motor",
            "observaciones",
            "activo",
            "creado_en",
            "actualizado_en",
        )
        .order_by("marca", "modelo", "pk")
    )
    for moto in motos:
        yield (
            moto.pk,
            moto.patente,
            moto.marca,
            moto.modelo,
            moto.anio,
            moto.cilindrada_cc,
            moto.color,
            moto.kilometraje_actual,
            moto.cliente_id,
            moto.cliente.nombre_completo,
            moto.cliente.telefono,
            moto.numero_chasis,
            moto.numero_motor,
            moto.observaciones,
            moto.activo,
            moto.creado_en,
            moto.actualizado_en,
        )


def filas_servicios():
    servicios = (
        Servicio.objects.select_related("moto", "cliente", "creado_por")
        .only(
            "pk",
            "fecha",
            "estado",
            "moto_id",
            "moto__patente",
            "moto__marca",
            "moto__modelo",
            "cliente_id",
            "cliente__nombre",
            "cliente__apellido",
            "kilometraje",
            "precio_total",
            "trabajos_adicionales",
            "observaciones",
            "creado_por__username",
            "creado_en",
            "actualizado_en",
        )
        .prefetch_related(
            Prefetch(
                "mantenimientos_realizados",
                queryset=MantenimientoRealizado.objects.select_related(
                    "tipo_mantenimiento"
                ).only(
                    "servicio_id",
                    "tipo_mantenimiento__nombre",
                ),
            )
        )
        .order_by("fecha", "pk")
    )
    for servicio in servicios:
        mantenimientos = " | ".join(
            realizado.tipo_mantenimiento.nombre
            for realizado in servicio.mantenimientos_realizados.all()
        )
        yield (
            servicio.pk,
            servicio.fecha,
            servicio.get_estado_display(),
            servicio.moto_id,
            servicio.moto.patente,
            servicio.moto.marca,
            servicio.moto.modelo,
            servicio.cliente_id,
            servicio.cliente.nombre_completo,
            servicio.kilometraje,
            servicio.precio_total,
            mantenimientos,
            servicio.trabajos_adicionales,
            servicio.observaciones,
            _nombre_usuario(servicio.creado_por),
            servicio.creado_en,
            servicio.actualizado_en,
        )


def filas_mantenimientos():
    mantenimientos = (
        MantenimientoRealizado.objects.select_related(
            "servicio__moto",
            "servicio__cliente",
            "tipo_mantenimiento",
        )
        .only(
            "pk",
            "servicio_id",
            "servicio__fecha",
            "servicio__estado",
            "servicio__moto_id",
            "servicio__moto__patente",
            "servicio__moto__marca",
            "servicio__moto__modelo",
            "servicio__cliente_id",
            "servicio__cliente__nombre",
            "servicio__cliente__apellido",
            "tipo_mantenimiento_id",
            "tipo_mantenimiento__nombre",
            "servicio__kilometraje",
            "observaciones",
            "creado_en",
        )
        .order_by("servicio__fecha", "servicio_id", "pk")
    )
    for realizado in mantenimientos:
        servicio = realizado.servicio
        yield (
            realizado.pk,
            servicio.pk,
            servicio.fecha,
            servicio.get_estado_display(),
            servicio.moto_id,
            servicio.moto.patente,
            servicio.moto.marca,
            servicio.moto.modelo,
            servicio.cliente_id,
            servicio.cliente.nombre_completo,
            realizado.tipo_mantenimiento_id,
            realizado.tipo_mantenimiento.nombre,
            servicio.kilometraje,
            realizado.observaciones,
            realizado.creado_en,
        )


def filas_tipos_mantenimiento():
    tipos = TipoMantenimiento.objects.order_by("nombre", "pk")
    for tipo in tipos:
        yield (
            tipo.pk,
            tipo.nombre,
            tipo.descripcion,
            tipo.activo,
            tipo.genera_recordatorio,
            tipo.intervalo_meses,
            tipo.intervalo_km,
            tipo.aviso_dias,
            tipo.aviso_km,
            tipo.creado_en,
            tipo.actualizado_en,
        )


def filas_seguimientos():
    seguimientos = (
        SeguimientoMantenimiento.objects.select_related(
            "mantenimiento_base__tipo_mantenimiento",
            "mantenimiento_base__servicio__moto",
            "cliente",
            "ultimo_contacto_por",
            "actualizado_por",
        )
        .only(
            "pk",
            "mantenimiento_base_id",
            "mantenimiento_base__tipo_mantenimiento_id",
            "mantenimiento_base__tipo_mantenimiento__nombre",
            "mantenimiento_base__servicio__moto_id",
            "mantenimiento_base__servicio__moto__patente",
            "cliente_id",
            "cliente__nombre",
            "cliente__apellido",
            "estado",
            "pospuesto_hasta",
            "turno_para",
            "ultimo_contacto_en",
            "ultimo_contacto_por__username",
            "observaciones",
            "actualizado_por__username",
            "creado_en",
            "actualizado_en",
        )
        .order_by("creado_en", "pk")
    )
    for seguimiento in seguimientos:
        mantenimiento = seguimiento.mantenimiento_base
        moto = mantenimiento.servicio.moto
        yield (
            seguimiento.pk,
            seguimiento.mantenimiento_base_id,
            mantenimiento.tipo_mantenimiento_id,
            mantenimiento.tipo_mantenimiento.nombre,
            moto.pk,
            moto.patente,
            seguimiento.cliente_id,
            seguimiento.cliente.nombre_completo,
            seguimiento.get_estado_display(),
            seguimiento.pospuesto_hasta,
            seguimiento.turno_para,
            seguimiento.ultimo_contacto_en,
            _nombre_usuario(seguimiento.ultimo_contacto_por),
            seguimiento.observaciones,
            _nombre_usuario(seguimiento.actualizado_por),
            seguimiento.creado_en,
            seguimiento.actualizado_en,
        )


def filas_eventos_seguimiento():
    eventos = (
        EventoSeguimientoMantenimiento.objects.select_related(
            "seguimiento__mantenimiento_base__tipo_mantenimiento",
            "seguimiento__mantenimiento_base__servicio__moto",
            "seguimiento__cliente",
            "usuario",
        )
        .only(
            "pk",
            "seguimiento_id",
            "tipo_evento",
            "usuario__username",
            "nota",
            "pospuesto_hasta",
            "turno_para",
            "creado_en",
            "seguimiento__mantenimiento_base__tipo_mantenimiento_id",
            "seguimiento__mantenimiento_base__tipo_mantenimiento__nombre",
            "seguimiento__mantenimiento_base__servicio__moto_id",
            "seguimiento__mantenimiento_base__servicio__moto__patente",
            "seguimiento__cliente_id",
            "seguimiento__cliente__nombre",
            "seguimiento__cliente__apellido",
        )
        .order_by("creado_en", "pk")
    )
    for evento in eventos:
        seguimiento = evento.seguimiento
        mantenimiento = seguimiento.mantenimiento_base
        moto = mantenimiento.servicio.moto
        yield (
            evento.pk,
            seguimiento.pk,
            evento.get_tipo_evento_display(),
            _nombre_usuario(evento.usuario),
            evento.nota,
            evento.pospuesto_hasta,
            evento.turno_para,
            evento.creado_en,
            mantenimiento.tipo_mantenimiento_id,
            mantenimiento.tipo_mantenimiento.nombre,
            moto.pk,
            moto.patente,
            seguimiento.cliente_id,
            seguimiento.cliente.nombre_completo,
        )


TABLAS_EXPORTACION = {
    "clientes": TablaExportacion(
        "clientes",
        "Clientes",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Nombre", 20),
            ColumnaExportacion("Apellido", 20),
            ColumnaExportacion("Teléfono", 18),
            ColumnaExportacion("Email", 28),
            ColumnaExportacion("Dirección", 32),
            ColumnaExportacion("Observaciones", 45, envolver=True),
            ColumnaExportacion("Activo", 10),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Actualizado", 18, FORMATO_FECHA_HORA),
        ),
        filas_clientes,
    ),
    "motos": TablaExportacion(
        "motos",
        "Motos",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Patente", 14),
            ColumnaExportacion("Marca", 18),
            ColumnaExportacion("Modelo", 22),
            ColumnaExportacion("Año", 10),
            ColumnaExportacion("Cilindrada cc", 14, FORMATO_ENTERO),
            ColumnaExportacion("Color", 16),
            ColumnaExportacion("Último kilometraje registrado", 24, FORMATO_ENTERO),
            ColumnaExportacion("Cliente actual ID", 16),
            ColumnaExportacion("Cliente actual", 26),
            ColumnaExportacion("Teléfono cliente actual", 22),
            ColumnaExportacion("Número chasis", 24),
            ColumnaExportacion("Número motor", 24),
            ColumnaExportacion("Observaciones", 45, envolver=True),
            ColumnaExportacion("Activa", 10),
            ColumnaExportacion("Creada", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Actualizada", 18, FORMATO_FECHA_HORA),
        ),
        filas_motos,
    ),
    "servicios": TablaExportacion(
        "servicios",
        "Servicios",
        (
            ColumnaExportacion("ID servicio", 12),
            ColumnaExportacion("Fecha", 14, FORMATO_FECHA),
            ColumnaExportacion("Estado", 14),
            ColumnaExportacion("Moto ID", 10),
            ColumnaExportacion("Patente", 14),
            ColumnaExportacion("Marca", 18),
            ColumnaExportacion("Modelo", 22),
            ColumnaExportacion("Cliente histórico ID", 18),
            ColumnaExportacion("Cliente histórico", 26),
            ColumnaExportacion("Kilometraje", 16, FORMATO_ENTERO),
            ColumnaExportacion("Precio total", 16, FORMATO_DINERO),
            ColumnaExportacion("Mantenimientos realizados", 35, envolver=True),
            ColumnaExportacion("Trabajos adicionales", 45, envolver=True),
            ColumnaExportacion("Observaciones", 45, envolver=True),
            ColumnaExportacion("Creado por", 18),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Actualizado", 18, FORMATO_FECHA_HORA),
        ),
        filas_servicios,
    ),
    "mantenimientos": TablaExportacion(
        "mantenimientos",
        "Mantenimientos",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Servicio ID", 12),
            ColumnaExportacion("Fecha servicio", 16, FORMATO_FECHA),
            ColumnaExportacion("Estado servicio", 16),
            ColumnaExportacion("Moto ID", 10),
            ColumnaExportacion("Patente", 14),
            ColumnaExportacion("Marca", 18),
            ColumnaExportacion("Modelo", 22),
            ColumnaExportacion("Cliente histórico ID", 18),
            ColumnaExportacion("Cliente histórico", 26),
            ColumnaExportacion("Tipo mantenimiento ID", 18),
            ColumnaExportacion("Tipo mantenimiento", 28),
            ColumnaExportacion("Kilometraje servicio", 20, FORMATO_ENTERO),
            ColumnaExportacion("Observaciones mantenimiento", 45, envolver=True),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
        ),
        filas_mantenimientos,
    ),
    "tipos-mantenimiento": TablaExportacion(
        "tipos-mantenimiento",
        "Tipos de mantenimiento",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Nombre", 28),
            ColumnaExportacion("Descripción", 45, envolver=True),
            ColumnaExportacion("Activo", 10),
            ColumnaExportacion("Genera recordatorio", 20),
            ColumnaExportacion("Intervalo meses", 16, FORMATO_ENTERO),
            ColumnaExportacion("Intervalo km", 16, FORMATO_ENTERO),
            ColumnaExportacion("Aviso días", 14, FORMATO_ENTERO),
            ColumnaExportacion("Aviso km", 14, FORMATO_ENTERO),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Actualizado", 18, FORMATO_FECHA_HORA),
        ),
        filas_tipos_mantenimiento,
    ),
    "seguimientos": TablaExportacion(
        "seguimientos",
        "Seguimientos",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Mantenimiento base ID", 20),
            ColumnaExportacion("Tipo mantenimiento ID", 18),
            ColumnaExportacion("Tipo mantenimiento", 28),
            ColumnaExportacion("Moto ID", 10),
            ColumnaExportacion("Patente", 14),
            ColumnaExportacion("Cliente contactado ID", 20),
            ColumnaExportacion("Cliente contactado", 26),
            ColumnaExportacion("Estado seguimiento", 20),
            ColumnaExportacion("Pospuesto hasta", 16, FORMATO_FECHA),
            ColumnaExportacion("Turno para", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Último contacto", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Último contacto por", 20),
            ColumnaExportacion("Observaciones", 45, envolver=True),
            ColumnaExportacion("Actualizado por", 18),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Actualizado", 18, FORMATO_FECHA_HORA),
        ),
        filas_seguimientos,
    ),
    "eventos-seguimiento": TablaExportacion(
        "eventos-seguimiento",
        "Eventos seguimiento",
        (
            ColumnaExportacion("ID", 10),
            ColumnaExportacion("Seguimiento ID", 14),
            ColumnaExportacion("Tipo evento", 18),
            ColumnaExportacion("Usuario", 18),
            ColumnaExportacion("Nota", 45, envolver=True),
            ColumnaExportacion("Pospuesto hasta", 16, FORMATO_FECHA),
            ColumnaExportacion("Turno para", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Creado", 18, FORMATO_FECHA_HORA),
            ColumnaExportacion("Tipo mantenimiento ID", 18),
            ColumnaExportacion("Tipo mantenimiento", 28),
            ColumnaExportacion("Moto ID", 10),
            ColumnaExportacion("Patente", 14),
            ColumnaExportacion("Cliente contactado ID", 20),
            ColumnaExportacion("Cliente contactado", 26),
        ),
        filas_eventos_seguimiento,
    ),
}


def crear_csv(clave):
    tabla = TABLAS_EXPORTACION[clave]
    salida = StringIO(newline="")
    salida.write("\ufeff")
    escritor = csv.writer(salida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(tabla.encabezados)
    for fila in tabla.obtener_filas():
        escritor.writerow(valor_csv(valor) for valor in fila)
    return salida.getvalue().encode("utf-8")


def _escribir_tabla(workbook, tabla):
    hoja = workbook.create_sheet(tabla.hoja)
    hoja.sheet_view.showGridLines = False
    hoja.freeze_panes = "A2"
    hoja.append(tabla.encabezados)

    relleno = PatternFill("solid", fgColor="0D6EFD")
    for indice, columna in enumerate(tabla.columnas, start=1):
        celda = hoja.cell(row=1, column=indice)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = columna.ancho

    cantidad = 0
    for fila in tabla.obtener_filas():
        hoja.append(tuple(valor_excel(valor) for valor in fila))
        cantidad += 1
        fila_excel = hoja.max_row
        for indice, columna in enumerate(tabla.columnas, start=1):
            celda = hoja.cell(row=fila_excel, column=indice)
            if columna.formato and celda.value != "":
                celda.number_format = columna.formato
            celda.alignment = Alignment(
                vertical="top",
                wrap_text=columna.envolver,
            )

    ultima_columna = get_column_letter(len(tabla.columnas))
    hoja.auto_filter.ref = f"A1:{ultima_columna}{hoja.max_row}"
    return cantidad


def _escribir_resumen(hoja, usuario, generado_en, cantidades):
    hoja.sheet_view.showGridLines = False
    hoja.column_dimensions["A"].width = 30
    hoja.column_dimensions["B"].width = 34
    hoja["A1"] = "MotoService"
    hoja["A1"].font = Font(size=18, bold=True, color="0D6EFD")
    hoja["A3"] = "Exportación generada"
    hoja["B3"] = valor_excel(generado_en)
    hoja["B3"].number_format = FORMATO_FECHA_HORA
    hoja["A4"] = "Generada por"
    hoja["B4"] = valor_excel(_nombre_usuario(usuario))

    resumen = (
        ("Clientes", "clientes"),
        ("Motos", "motos"),
        ("Servicios", "servicios"),
        ("Mantenimientos realizados", "mantenimientos"),
        ("Tipos de mantenimiento", "tipos-mantenimiento"),
        ("Seguimientos", "seguimientos"),
        ("Eventos", "eventos-seguimiento"),
    )
    for fila, (etiqueta, clave) in enumerate(resumen, start=6):
        hoja.cell(row=fila, column=1, value=etiqueta)
        hoja.cell(row=fila, column=2, value=cantidades[clave])

    fila_advertencia = 15
    hoja.merge_cells(
        start_row=fila_advertencia,
        start_column=1,
        end_row=fila_advertencia,
        end_column=2,
    )
    advertencia = hoja.cell(row=fila_advertencia, column=1)
    advertencia.value = (
        "Este archivo es una exportación de datos y no reemplaza un backup "
        "PostgreSQL."
    )
    advertencia.font = Font(bold=True, color="842029")
    advertencia.alignment = Alignment(wrap_text=True)


def crear_workbook_completo(usuario, *, generado_en=None):
    generado_en = momento_local(generado_en)
    workbook = Workbook()
    workbook.iso_dates = True
    workbook.properties.creator = "MotoService"
    workbook.properties.title = "Exportación completa de datos operativos"
    resumen = workbook.active
    resumen.title = "Resumen"

    cantidades = {}
    for clave, tabla in TABLAS_EXPORTACION.items():
        cantidades[clave] = _escribir_tabla(workbook, tabla)

    _escribir_resumen(resumen, usuario, generado_en, cantidades)
    return workbook


def crear_excel_completo(usuario, *, generado_en=None):
    workbook = crear_workbook_completo(usuario, generado_en=generado_en)
    salida = BytesIO()
    workbook.save(salida)
    workbook.close()
    return salida.getvalue()
