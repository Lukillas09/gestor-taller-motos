from datetime import date
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from apps.clientes.models import Cliente
from apps.exportaciones.services import crear_excel_completo
from apps.mantenimientos.models import MantenimientoRealizado, TipoMantenimiento
from apps.motos.models import Moto
from apps.notificaciones.models import (
    EventoSeguimientoMantenimiento,
    SeguimientoMantenimiento,
)
from apps.servicios.models import Servicio


class ExcelCompletoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.usuario = get_user_model().objects.create_user(
            username="lucas",
            password="clave-segura",
        )
        cls.carlos = Cliente.objects.create(nombre="Carlos", apellido="González")
        cls.martin = Cliente.objects.create(nombre="Martín", apellido="Pérez")
        Cliente.objects.create(
            nombre="=2+2",
            observaciones="@SUM(A1:A2)",
        )
        cls.moto = Moto.objects.create(
            cliente=cls.carlos,
            patente="AF123XY",
            marca="Honda",
            modelo="Tornado",
            kilometraje_actual=15000,
            observaciones="Observación extensa " * 30,
        )
        Moto.objects.create(
            cliente=cls.carlos,
            marca="+cmd",
            modelo="Modelo seguro",
        )
        cls.servicio = Servicio.objects.create(
            moto=cls.moto,
            fecha=date(2026, 9, 1),
            kilometraje=14000,
            precio_total=Decimal("75000.10"),
            observaciones="@SUM(A1:A2)",
            creado_por=cls.usuario,
        )
        cls.tipo = TipoMantenimiento.objects.get(nombre="Cambio de aceite")
        cls.tipo_peligroso = TipoMantenimiento.objects.create(nombre="-Trabajo")
        cls.realizado = MantenimientoRealizado.objects.create(
            servicio=cls.servicio,
            tipo_mantenimiento=cls.tipo,
        )
        MantenimientoRealizado.objects.create(
            servicio=cls.servicio,
            tipo_mantenimiento=cls.tipo_peligroso,
        )
        cls.turno_para = timezone.now().replace(microsecond=0)
        cls.seguimiento = SeguimientoMantenimiento.objects.create(
            mantenimiento_base=cls.realizado,
            cliente=cls.carlos,
            estado=SeguimientoMantenimiento.Estado.TURNO_ACORDADO,
            turno_para=cls.turno_para,
            actualizado_por=cls.usuario,
            observaciones="=Nota peligrosa",
        )
        cls.evento = EventoSeguimientoMantenimiento.objects.create(
            seguimiento=cls.seguimiento,
            tipo_evento=EventoSeguimientoMantenimiento.Tipo.TURNO_ACORDADO,
            usuario=cls.usuario,
            nota="+Otra nota",
            turno_para=cls.turno_para,
        )
        cls.moto.cliente = cls.martin
        cls.moto.save()

    def _workbook(self):
        contenido = crear_excel_completo(self.usuario)
        return load_workbook(BytesIO(contenido), data_only=False)

    def _fila_por_id(self, hoja, identificador):
        encabezados = [celda.value for celda in hoja[1]]
        for valores in hoja.iter_rows(min_row=2, values_only=True):
            if valores[0] == identificador:
                return dict(zip(encabezados, valores, strict=True))
        self.fail(f"No se encontró la fila {identificador} en {hoja.title}.")

    def test_excel_incluye_todas_las_hojas_requeridas(self):
        workbook = self._workbook()

        self.assertEqual(
            workbook.sheetnames,
            [
                "Resumen",
                "Clientes",
                "Motos",
                "Servicios",
                "Mantenimientos",
                "Tipos de mantenimiento",
                "Seguimientos",
                "Eventos seguimiento",
            ],
        )

    def test_excel_conserva_relaciones_historicas(self):
        workbook = self._workbook()

        servicio = self._fila_por_id(workbook["Servicios"], self.servicio.pk)
        moto = self._fila_por_id(workbook["Motos"], self.moto.pk)
        seguimiento = self._fila_por_id(
            workbook["Seguimientos"],
            self.seguimiento.pk,
        )
        self.assertEqual(servicio["Cliente histórico"], "Carlos González")
        self.assertEqual(moto["Cliente actual"], "Martín Pérez")
        self.assertEqual(seguimiento["Cliente contactado"], "Carlos González")

    def test_excel_no_contiene_formulas_inyectadas(self):
        workbook = self._workbook()

        for hoja in workbook.worksheets:
            for fila in hoja.iter_rows():
                for celda in fila:
                    self.assertNotEqual(celda.data_type, "f")

        clientes = workbook["Clientes"]
        valores_nombres = [celda.value for celda in clientes["B"]]
        self.assertIn("'=2+2", valores_nombres)
        servicio = self._fila_por_id(workbook["Servicios"], self.servicio.pk)
        self.assertEqual(servicio["Observaciones"], "'@SUM(A1:A2)")

    def test_excel_convierte_datetime_a_hora_local_naive(self):
        workbook = self._workbook()

        seguimiento = self._fila_por_id(
            workbook["Seguimientos"],
            self.seguimiento.pk,
        )
        esperado = timezone.localtime(self.turno_para).replace(tzinfo=None)
        self.assertEqual(seguimiento["Turno para"], esperado)
        self.assertIsNone(seguimiento["Turno para"].tzinfo)

    def test_excel_aplica_formato_legible_sin_expandir_textos_largos(self):
        workbook = self._workbook()
        servicios = workbook["Servicios"]
        motos = workbook["Motos"]

        self.assertEqual(servicios.freeze_panes, "A2")
        self.assertTrue(servicios.auto_filter.ref)
        self.assertTrue(servicios["A1"].font.bold)
        encabezados_servicios = [celda.value for celda in servicios[1]]
        indice_precio = encabezados_servicios.index("Precio total") + 1
        fila_servicio = next(
            fila
            for fila in range(2, servicios.max_row + 1)
            if servicios.cell(fila, 1).value == self.servicio.pk
        )
        celda_precio = servicios.cell(fila_servicio, indice_precio)
        self.assertEqual(celda_precio.value, 75000.1)
        self.assertIn("0.00", celda_precio.number_format)

        encabezados_motos = [celda.value for celda in motos[1]]
        indice_observaciones = encabezados_motos.index("Observaciones") + 1
        letra = get_column_letter(indice_observaciones)
        self.assertLessEqual(motos.column_dimensions[letra].width, 45)
        fila_moto = next(
            fila
            for fila in range(2, motos.max_row + 1)
            if motos.cell(fila, 1).value == self.moto.pk
        )
        self.assertTrue(motos.cell(fila_moto, indice_observaciones).alignment.wrap_text)

    def test_resumen_contiene_conteos_y_advertencia_sin_secrets(self):
        workbook = self._workbook()
        resumen = workbook["Resumen"]
        contenido = " ".join(
            str(celda.value)
            for fila in resumen.iter_rows()
            for celda in fila
            if celda.value is not None
        )

        self.assertIn("lucas", contenido)
        self.assertIn("no reemplaza un backup PostgreSQL", contenido)
        self.assertNotIn("DATABASE_URL", contenido)
        self.assertNotIn("SECRET_KEY", contenido)

    def test_excel_completo_usa_cantidad_constante_de_consultas(self):
        with self.assertNumQueries(8):
            contenido = crear_excel_completo(self.usuario)

        self.assertTrue(contenido.startswith(b"PK"))
